"""Read a .docx file into a SourceDoc: one addressable paragraph id per paragraph, tables
as plain rows, and the document's own core-properties date (T7).

The metadata date is never a value-selecting date (R3): it exists only so the meeting-date
fallback (DESIGN.md section 4.1) can derive a tax year when no date is stated in the text.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from agent_pipeline.sources.document import SourceDoc


def read_docx(path: Path) -> SourceDoc:
    doc = Document(str(path))

    paragraphs: dict[str, str] = {}
    for i, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if text:
            paragraphs[f"p{i}"] = text

    tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append([row for row in rows if any(row)])

    metadata_date = doc.core_properties.modified or doc.core_properties.created
    date_only = metadata_date.date() if metadata_date is not None else None

    return SourceDoc(path=path, paragraphs=paragraphs, tables=tables, metadata_date=date_only)
