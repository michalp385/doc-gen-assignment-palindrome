"""Writes a synthetic client's folder from a scenario (T27, D4, DESIGN.md section 10.8): the
account JSON, the meeting note and the general document (python-docx), the report-instruction
table (python-docx), the statement image (Pillow), the internal-guidance notes and the
report spec -- the same files a real client folder has.

Deterministic: every .docx gets a fixed core-properties date and normalised zip timestamps, so
writing the same scenario and note twice gives byte-identical files (the frozen clients under
`data/synthetic/generated/` are written once and never regenerated per run). The docx helpers
mirror `scripts/build_handwritten.py`'s, which is a script and so not importable.
"""

from __future__ import annotations

import json
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from PIL import Image, ImageDraw, ImageFont

from report_eval.synth.scenario import (
    Scenario,
    account_data,
    distractor_text,
    instruction_fields,
)

SPEC_SOURCE = Path("data/client_01_clean/template_spec.md")

_GUIDANCE = (
    "# Internal notes: data sources\n\n"
    "For whoever configures the report: what each source is, and how they relate.\n\n"
    "## The sources\n\n"
    "- **client_data_db.json**: a snapshot of the client's accounts from the custody systems. "
    "It is the record of which accounts exist and who holds them; each account carries its own "
    "valuation date.\n"
    "- **meeting_notes.docx**: the adviser's file note of the latest review meeting.\n"
    "- **report_request.docx**: the adviser's instruction for this report.\n"
)


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in text.split("\n\n") if p.strip()]


def _freeze_zip_timestamps(path: Path, when: datetime) -> None:
    date_time = (when.year, when.month, when.day, 0, 0, 0)
    with zipfile.ZipFile(path, "r") as src:
        entries = [(info, src.read(info.filename)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, data in entries:
            info.date_time = date_time
            dst.writestr(info, data)


def _new_document(when: datetime) -> DocumentType:
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.created = when
    doc.core_properties.modified = when
    return doc


def _write_paragraphs(path: Path, paragraphs: list[str], when: datetime) -> None:
    doc = _new_document(when)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(str(path))
    _freeze_zip_timestamps(path, when)


def _write_instruction(path: Path, fields: list[tuple[str, str]], when: datetime) -> None:
    doc = _new_document(when)
    doc.add_paragraph("Report Requirement Summary")
    table = doc.add_table(rows=0, cols=2)
    for label, value in fields:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
    doc.save(str(path))
    _freeze_zip_timestamps(path, when)


def _write_statement_image(path: Path, title: str, rows: list[tuple[str, str, str, str]]) -> None:
    header = ("Account", "Type", "Value", "Valued on")
    col_widths = [220, 220, 110, 110]
    row_height = 36
    top = 90
    width = sum(col_widths) + 20
    height = top + row_height * (len(rows) + 1) + 20
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    # The built-in font has no £ glyph, so statement cells write amounts as "GBP 1,234".
    font = ImageFont.load_default(size=16)
    draw.text((10, 10), title, fill="black", font=font)

    def draw_row(y: int, cells: tuple[str, ...], *, header_row: bool) -> None:
        x = 10
        if header_row:
            draw.rectangle([x, y, x + sum(col_widths), y + row_height], fill=(230, 230, 230))
        for cell, w in zip(cells, col_widths, strict=True):
            draw.rectangle([x, y, x + w, y + row_height], outline="black")
            draw.text((x + 5, y + row_height // 2 - 5), cell, fill="black", font=font)
            x += w

    draw_row(top, header, header_row=True)
    for i, row in enumerate(rows):
        draw_row(top + row_height * (i + 1), row, header_row=False)
    image.save(path, format="PNG")


def write_client(scenario: Scenario, meeting_note: str, out_dir: Path) -> None:
    """Write the folder for one scenario. `meeting_note` is the already-checked prose
    (`prose.generate_prose`); nothing here calls a model."""
    out_dir.mkdir(parents=True, exist_ok=True)
    when = datetime(
        scenario.meeting_date.year,
        scenario.meeting_date.month,
        scenario.meeting_date.day,
        tzinfo=timezone.utc,
    )

    (out_dir / "client_data_db.json").write_text(
        json.dumps(account_data(scenario), indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    _write_paragraphs(out_dir / "meeting_notes.docx", _paragraphs(meeting_note), when)
    _write_instruction(out_dir / "report_request.docx", instruction_fields(scenario), when)
    _write_paragraphs(
        out_dir / "platform_market_update.docx", _paragraphs(distractor_text(scenario)), when
    )
    _write_statement_image(
        out_dir / "statement_summary.png",
        f"{scenario.platform} Account Statement",
        [
            (
                account.id,
                account.type,
                f"GBP {int(account.value):,}" if account.value is not None else "",
                account.valuation_date.isoformat() if account.valuation_date else "",
            )
            for account in scenario.accounts
            if account.status == "open"
        ],
    )
    (out_dir / "fde_notes.md").write_text(_GUIDANCE, encoding="utf-8")
    shutil.copyfile(SPEC_SOURCE, out_dir / "template_spec.md")
