"""Leading and trailing apostrophes and quotes in the overlap tokenisers (T29 verifier
checkpoint, second pass), tests first.

Folding a curly quote into a straight apostrophe made `‘cash` and `clients’` tokens that no
plain-text prompt contains, which splits a shared six-word run just as the curly internal
apostrophe did. A token keeps an apostrophe only inside a word ("client's").
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path
from types import ModuleType

import pytest
from docx import Document

from report_eval.synth import phrases

ROOT = Path(__file__).resolve().parent.parent


def _load_check_repo() -> ModuleType:
    spec = importlib.util.spec_from_file_location("check_repo_tok2", ROOT / "scripts/check_repo.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


CHECK_REPO = _load_check_repo()

TOKENISERS = [CHECK_REPO._words, phrases._words]  # noqa: SLF001  # units under test


@pytest.mark.parametrize("words", TOKENISERS)
def test_a_quoted_word_tokenises_like_the_plain_word(words) -> None:  # type: ignore[no-untyped-def]  # parametrised callable
    assert words("see ‘cash’ account") == words("see cash account")


@pytest.mark.parametrize("words", TOKENISERS)
def test_a_plural_possessive_tokenises_like_the_plural(words) -> None:  # type: ignore[no-untyped-def]  # parametrised callable
    assert words("the clients’ accounts") == words("the clients accounts")
    assert words("the clients' accounts") == words("the clients accounts")


@pytest.mark.parametrize("words", TOKENISERS)
def test_an_internal_apostrophe_is_kept(words) -> None:  # type: ignore[no-untyped-def]  # parametrised callable
    assert words("the client’s account") == ["the", "client's", "account"]


def test_a_prompt_quoting_a_source_with_quoted_words_is_flagged(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(CHECK_REPO, "ROOT", tmp_path)
    monkeypatch.setattr(CHECK_REPO, "DATA", tmp_path / "data")
    monkeypatch.setattr(CHECK_REPO, "PROMPTS_DIR", tmp_path / "config" / "prompts")
    (tmp_path / "config" / "prompts").mkdir(parents=True)
    (tmp_path / "data" / "client_x").mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("alpha so ‘cash’ for the clients’ accounts omega")
    doc.save(str(tmp_path / "data" / "client_x" / "meeting_notes.docx"))
    (tmp_path / "config" / "prompts" / "p.md").write_text(
        "beta so cash for the clients accounts gamma\n", encoding="utf-8"
    )
    assert CHECK_REPO.check_prompt_overlap({})


def test_unrelated_text_is_not_flagged(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(CHECK_REPO, "ROOT", tmp_path)
    monkeypatch.setattr(CHECK_REPO, "DATA", tmp_path / "data")
    monkeypatch.setattr(CHECK_REPO, "PROMPTS_DIR", tmp_path / "config" / "prompts")
    (tmp_path / "config" / "prompts").mkdir(parents=True)
    (tmp_path / "data" / "client_x").mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("A completely different note about something else entirely.")
    doc.save(str(tmp_path / "data" / "client_x" / "meeting_notes.docx"))
    (tmp_path / "config" / "prompts" / "p.md").write_text(
        "Nothing shared here.\n", encoding="utf-8"
    )
    assert CHECK_REPO.check_prompt_overlap({}) == []
