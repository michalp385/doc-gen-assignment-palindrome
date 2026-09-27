"""Source classification (T12): structural checks, model routing, quote verification,
confidence threshold, and the consistency/stop rules (DESIGN.md sections 3.1-3.2, 8.4).
A stub classifier stands in for the real model call -- offline, no network.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path

import pytest
from docx import Document

from agent_pipeline.sources.classify import (
    ClassificationStopError,
    ClassificationVerdict,
    classify,
)


@dataclass
class StubClassifier:
    # (marker, verdict): `marker` must be a substring of the fixture's own real content, so
    # routing never depends on the verdict's own evidence_quote -- some tests deliberately
    # give a verdict whose quote doesn't match the file, to test that path.
    entries: list[tuple[str, ClassificationVerdict]]
    calls: list[str] = field(default_factory=list)

    def classify_text(self, text: str) -> ClassificationVerdict:
        self.calls.append(text)
        for marker, verdict in self.entries:
            if marker in text:
                return verdict
        raise AssertionError(f"no stub entry matches text: {text[:80]!r}")


def _account_data_json(value: int = 10000) -> str:
    return json.dumps(
        {
            "holders": {
                "client": {
                    "name": "Test Person",
                    "accounts": [
                        {
                            "account_id": "T-1",
                            "type": "Cash Account",
                            "owner": "Test Person",
                            "status": "open",
                            "value": value,
                            "currency": "GBP",
                        }
                    ],
                }
            }
        }
    )


def _write(tmp_path: Path, name: str, content: str) -> Path:
    path = tmp_path / name
    path.write_text(content, encoding="utf-8")
    return path


# Real fixture content, reused across tests so a stub's marker always matches the actual
# file, independent of whatever a test's verdict claims.
_INSTRUCTION_TEXT = "# Report Requirement Summary\n\nCover the ISA."
_MEETING_TEXT = "# Meeting note\n\nWe discussed the ISA top-up."


def _base_folder(tmp_path: Path) -> Path:
    _write(tmp_path, "accounts.json", _account_data_json())
    _write(tmp_path, "instruction.md", _INSTRUCTION_TEXT)
    _write(tmp_path, "meeting.md", _MEETING_TEXT)
    return tmp_path


def _verdict(role: str, quote: str, confidence: float = 0.9) -> ClassificationVerdict:
    return ClassificationVerdict(role=role, evidence_quote=quote, confidence=confidence)  # type: ignore[arg-type]


def _default_classifier() -> StubClassifier:
    return StubClassifier(
        [
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
        ]
    )


# --- Structural classification -------------------------------------------------------------


def test_valid_account_data_json_is_classified_structurally(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    result = classify(folder, _default_classifier())

    accounts = [s for s in result.sources if s.role == "account_data"]
    assert len(accounts) == 1
    assert accounts[0].method == "structural"


def test_invalid_json_becomes_unknown(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    _write(folder, "broken.json", "{ not valid json")
    result = classify(folder, _default_classifier())

    broken = next(s for s in result.sources if s.path.name == "broken.json")
    assert broken.role == "unknown"
    assert broken.method == "structural"


def test_image_file_is_a_statement_image_candidate(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    (folder / "statement.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    result = classify(folder, _default_classifier())

    image = next(s for s in result.sources if s.path.name == "statement.png")
    assert image.role == "statement_image"
    assert image.method == "structural"


def test_unrecognised_extension_is_unknown(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    _write(folder, "notes.pdf", "not readable by any adapter")
    result = classify(folder, _default_classifier())

    pdf = next(s for s in result.sources if s.path.name == "notes.pdf")
    assert pdf.role == "unknown"
    assert pdf.method == "structural"


# --- Model routing, quote verification, confidence -----------------------------------------


def test_model_classified_role_is_used_when_confident_and_verified(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    result = classify(folder, _default_classifier())

    instruction = next(s for s in result.sources if s.path.name == "instruction.md")
    assert instruction.role == "report_instruction"
    assert instruction.method == "model"
    assert instruction.confidence == 0.9


def test_a_docx_table_reaches_the_classifier_not_just_its_paragraphs(tmp_path: Path) -> None:
    # Regression: T12's live check found report_request.docx has exactly one paragraph (a
    # heading) and all its substance in a table; the classifier originally saw only the
    # heading and (correctly, given that little text) returned low confidence.
    folder = tmp_path
    _write(folder, "accounts.json", _account_data_json())
    _write(folder, "meeting.md", _MEETING_TEXT)

    doc = Document()
    doc.add_paragraph("Report Requirement Summary")
    table = doc.add_table(rows=0, cols=2)
    for label, value in [("Adviser", "Someone"), ("Accounts covered", "The ISA")]:
        row = table.add_row().cells
        row[0].text, row[1].text = label, value
    doc.save(str(folder / "instruction.docx"))

    calls: list[str] = []
    classifier = StubClassifier(
        [
            ("Accounts covered", _verdict("report_instruction", "Accounts covered | The ISA")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
        ],
        calls=calls,
    )
    classify(folder, classifier)

    assert any("Accounts covered" in call and "The ISA" in call for call in calls)


def test_unverified_evidence_quote_forces_unknown(tmp_path: Path) -> None:
    # A second, verified report_instruction keeps this out of stop-condition territory --
    # the point here is the unverified-quote file's own role, not the aggregate check.
    folder = _base_folder(tmp_path)
    _write(folder, "instruction_good.md", "# Report Requirement Summary\n\nCover the GIA.")
    classifier = StubClassifier(
        [
            ("Cover the ISA.", _verdict("report_instruction", "this text is not in the file")),
            ("Cover the GIA.", _verdict("report_instruction", "Cover the GIA.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
        ]
    )
    result = classify(folder, classifier)

    instruction = next(s for s in result.sources if s.path.name == "instruction.md")
    assert instruction.role == "unknown"
    assert "not found" in instruction.reason


def test_low_confidence_document_becomes_unknown_but_a_confident_one_still_satisfies_the_role(
    tmp_path: Path,
) -> None:
    # A single low-confidence candidate would correctly stop the run (zero *report_instruction*
    # sources remain, and that's a required role) -- so this needs a second, confident
    # candidate to isolate "low confidence -> unknown" from the separate stop-condition check.
    folder = tmp_path
    _write(folder, "accounts.json", _account_data_json())
    _write(folder, "instruction_bad.md", "# Report Requirement Summary\n\nWeak signal.")
    _write(folder, "instruction_good.md", _INSTRUCTION_TEXT)
    _write(folder, "meeting.md", _MEETING_TEXT)
    classifier = StubClassifier(
        [
            ("Weak signal.", _verdict("report_instruction", "Weak signal.", confidence=0.5)),
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
        ]
    )
    result = classify(folder, classifier, confidence_threshold=0.7)

    by_name = {s.path.name: s for s in result.sources}
    assert by_name["instruction_bad.md"].role == "unknown"
    assert "threshold" in by_name["instruction_bad.md"].reason
    assert by_name["instruction_good.md"].role == "report_instruction"


def test_confidence_exactly_at_the_threshold_is_accepted(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    classifier = StubClassifier(
        [
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.", confidence=0.7)),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
        ]
    )
    result = classify(folder, classifier, confidence_threshold=0.7)

    instruction = next(s for s in result.sources if s.path.name == "instruction.md")
    assert instruction.role == "report_instruction"


# --- Stop conditions (DESIGN.md section 8.4) ------------------------------------------------


def test_no_account_data_stops(tmp_path: Path) -> None:
    folder = tmp_path
    _write(folder, "instruction.md", _INSTRUCTION_TEXT)
    _write(folder, "meeting.md", _MEETING_TEXT)
    with pytest.raises(ClassificationStopError, match="account_data"):
        classify(folder, _default_classifier())


def test_no_report_instruction_stops(tmp_path: Path) -> None:
    folder = tmp_path
    _write(folder, "accounts.json", _account_data_json())
    _write(folder, "meeting.md", _MEETING_TEXT)
    classifier = StubClassifier(
        [
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            )
        ]
    )
    with pytest.raises(ClassificationStopError, match="report_instruction"):
        classify(folder, classifier)


def test_no_meeting_record_stops(tmp_path: Path) -> None:
    folder = tmp_path
    _write(folder, "accounts.json", _account_data_json())
    _write(folder, "instruction.md", _INSTRUCTION_TEXT)
    classifier = StubClassifier(
        [("Cover the ISA.", _verdict("report_instruction", "Cover the ISA."))]
    )
    with pytest.raises(ClassificationStopError, match="meeting_record"):
        classify(folder, classifier)


def test_two_disagreeing_report_instructions_stops(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    _write(folder, "instruction2.md", "# Report Requirement Summary\n\nCover the GIA instead.")
    classifier = StubClassifier(
        [
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.")),
            ("Cover the GIA instead.", _verdict("report_instruction", "Cover the GIA instead.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
        ]
    )
    with pytest.raises(ClassificationStopError, match="report_instruction"):
        classify(folder, classifier)


def test_two_disagreeing_account_data_sources_stops(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    _write(folder, "accounts2.json", _account_data_json(value=99999))
    with pytest.raises(ClassificationStopError, match="account_data"):
        classify(folder, _default_classifier())


# --- Deduplication ----------------------------------------------------------------------


def test_two_identical_report_instructions_are_deduplicated_with_a_note(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    _write(folder, "instruction_copy.md", _INSTRUCTION_TEXT)
    classifier = StubClassifier(
        [
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
        ]
    )
    result = classify(folder, classifier)

    instructions = [s for s in result.sources if s.role == "report_instruction"]
    assert len(instructions) == 1
    assert any("report_instruction" in note for note in result.notes)


def test_two_identical_account_data_files_are_deduplicated_with_a_note(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    _write(folder, "accounts_copy.json", _account_data_json())
    result = classify(folder, _default_classifier())

    accounts = [s for s in result.sources if s.role == "account_data"]
    assert len(accounts) == 1
    assert any("account_data" in note for note in result.notes)


def test_several_different_meeting_records_are_all_kept(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    _write(folder, "meeting2.md", "# Second meeting\n\nFollow-up call about charges.")
    classifier = StubClassifier(
        [
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
            (
                "Follow-up call about charges.",
                _verdict("meeting_record", "Follow-up call about charges."),
            ),
        ]
    )
    result = classify(folder, classifier)

    meetings = [s for s in result.sources if s.role == "meeting_record"]
    assert len(meetings) == 2
    assert result.notes == []


# --- General document / internal guidance route through the model like anything else -------


def test_general_document_and_internal_guidance_route_through_the_model(tmp_path: Path) -> None:
    folder = _base_folder(tmp_path)
    _write(folder, "market_update.md", "# Quarterly update\n\nGeneral market commentary.")
    _write(folder, "notes.md", "# Internal notes\n\nHow to configure this client's report.")
    classifier = StubClassifier(
        [
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
            (
                "General market commentary.",
                _verdict("general_document", "General market commentary."),
            ),
            (
                "How to configure this client's report.",
                _verdict("internal_guidance", "How to configure this client's report."),
            ),
        ]
    )
    result = classify(folder, classifier)

    by_name = {s.path.name: s for s in result.sources}
    assert by_name["market_update.md"].role == "general_document"
    assert by_name["notes.md"].role == "internal_guidance"


# --- The model returning a structural-only role (verifier report, T12 pre-commit checkpoint) -


def test_model_claiming_account_data_for_a_text_file_degrades_instead_of_crashing(
    tmp_path: Path,
) -> None:
    # Reproduces the verifier's finding: a text file the model mislabels "account_data" used
    # to reach _content_key's read_accounts() call on a non-JSON file and crash the whole
    # run. It must degrade to unknown instead, and never be counted as the required
    # account_data role (the real accounts.json still satisfies that on its own).
    folder = _base_folder(tmp_path)
    _write(folder, "confused.md", "# Some document\n\nAmbiguous content.")
    classifier = StubClassifier(
        [
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
            ("Ambiguous content.", _verdict("account_data", "Ambiguous content.")),
        ]
    )

    result = classify(folder, classifier)  # must not raise

    confused = next(s for s in result.sources if s.path.name == "confused.md")
    assert confused.role == "unknown"
    assert "structural-only" in confused.reason
    accounts = [s for s in result.sources if s.role == "account_data"]
    assert len(accounts) == 1
    assert accounts[0].path.name == "accounts.json"


def test_model_claiming_statement_image_for_a_text_file_degrades(tmp_path: Path) -> None:
    # A second, correctly-classified instruction file keeps this out of stop-condition
    # territory, so the per-file degrade itself is what's being checked (same pattern as the
    # low-confidence and unverified-quote tests above).
    folder = tmp_path
    _write(folder, "accounts.json", _account_data_json())
    _write(folder, "instruction_bad.md", "# Some document\n\nAmbiguous content.")
    _write(folder, "instruction_good.md", _INSTRUCTION_TEXT)
    _write(folder, "meeting.md", _MEETING_TEXT)
    classifier = StubClassifier(
        [
            ("Ambiguous content.", _verdict("statement_image", "Ambiguous content.")),
            ("Cover the ISA.", _verdict("report_instruction", "Cover the ISA.")),
            (
                "We discussed the ISA top-up.",
                _verdict("meeting_record", "We discussed the ISA top-up."),
            ),
        ]
    )

    result = classify(folder, classifier)  # must not raise

    bad = next(s for s in result.sources if s.path.name == "instruction_bad.md")
    assert bad.role == "unknown"
    assert "structural-only" in bad.reason
    good = next(s for s in result.sources if s.path.name == "instruction_good.md")
    assert good.role == "report_instruction"
