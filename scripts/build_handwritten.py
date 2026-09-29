"""Build the hand-written eval cases (T3) from their markdown sources.

Reads each `eval/handwritten_src/case_NN.md` and writes the client folder it describes to
`data/synthetic/handwritten/case_NN/`: an account-data JSON, one or more meeting-note and
report-instruction .docx files, and (where the case calls for them) an internal-guidance
.md file, an unknown .docx file, and a statement-image .png.

Deterministic: every .docx gets a fixed core-properties timestamp (or the case's own
`Metadata-Date:`, for the one case that exercises the metadata-date fallback), so rebuilding
reproduces the committed folder byte-for-byte (tests/test_handwritten_build.py checks this).

Usage:
    uv run python scripts/build_handwritten.py [case_NN ...] [--out-dir DIR]
        # default: every case, written under data/synthetic/handwritten/
"""

from __future__ import annotations

import json
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = ROOT / "eval" / "handwritten_src"
OUT_DIR = ROOT / "data" / "synthetic" / "handwritten"

# Every generated .docx gets this timestamp unless a section gives its own Metadata-Date.
DEFAULT_DOCX_DATE = datetime(2026, 1, 1, tzinfo=timezone.utc)

SECTION_RE = re.compile(r"^## (.+)$", re.MULTILINE)
DIRECTIVE_RE = re.compile(r"^(Filename|Metadata-Date):\s*(.+)$")
JSON_BLOCK_RE = re.compile(r"```json\s*\n(.*?)\n```", re.DOTALL)


@dataclass
class Section:
    heading: str
    directives: dict[str, str]
    body: str  # section text with any leading directive lines removed


def _split_sections(markdown: str) -> list[Section]:
    """Split a case file into its '## Heading' sections, each with leading directives parsed."""
    matches = list(SECTION_RE.finditer(markdown))
    sections = []
    for i, m in enumerate(matches):
        heading = m.group(1).strip()
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(markdown)
        body = markdown[start:end].strip("\n")
        directives: dict[str, str] = {}
        lines = body.split("\n")
        consumed = 0
        for line in lines:
            dm = DIRECTIVE_RE.match(line.strip())
            if dm is None:
                break
            directives[dm.group(1)] = dm.group(2).strip()
            consumed += 1
        body = "\n".join(lines[consumed:]).strip("\n")
        sections.append(Section(heading=heading, directives=directives, body=body))
    return sections


def _paragraphs(body: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", body.strip()) if p.strip()]


def _table_rows(body: str) -> list[tuple[str, ...]]:
    rows = []
    for line in body.strip().splitlines():
        line = line.strip()
        if not line:
            continue
        rows.append(tuple(part.strip() for part in line.split("|")))
    return rows


def _docx_date(section: Section) -> datetime:
    if "Metadata-Date" in section.directives:
        return datetime.fromisoformat(section.directives["Metadata-Date"]).replace(
            tzinfo=timezone.utc
        )
    return DEFAULT_DOCX_DATE


def _new_document(when: datetime) -> DocumentType:
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.created = when
    doc.core_properties.modified = when
    return doc


def _freeze_zip_timestamps(path: Path, when: datetime) -> None:
    """Rewrite every entry's zip-local timestamp to a fixed value.

    python-docx stamps each part with `time.localtime()` at save time (independent of
    `core_properties`), so two saves of identical content differ in raw bytes unless this
    is normalised afterwards. Needed for build_case's output to be reproducible (T3).
    """
    date_time = (when.year, when.month, when.day, 0, 0, 0)
    with zipfile.ZipFile(path, "r") as src:
        entries = [(info, src.read(info.filename)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, data in entries:
            info.date_time = date_time
            dst.writestr(info, data)


def _write_paragraphs_docx(path: Path, paragraphs: list[str], when: datetime) -> None:
    doc = _new_document(when)
    for para in paragraphs:
        doc.add_paragraph(para)
    doc.save(str(path))
    _freeze_zip_timestamps(path, when)


def _write_instruction_docx(path: Path, fields: list[tuple[str, str]], when: datetime) -> None:
    doc = _new_document(when)
    doc.add_paragraph("Report Requirement Summary")
    table = doc.add_table(rows=0, cols=2)
    for label, value in fields:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
    doc.save(str(path))
    _freeze_zip_timestamps(path, when)


def _render_statement_image(path: Path, title: str, rows: list[tuple[str, str, str, str]]) -> None:
    header = ("Account", "Type", "Value", "Valued on")
    col_widths = [220, 220, 110, 110]
    row_height = 36
    top = 90
    width = sum(col_widths) + 20
    height = top + row_height * (len(rows) + 1) + 20
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    # Pillow's built-in default font has no £ glyph even at a chosen size, so case sources
    # write amounts as "GBP 1,234" for statement-image cells (matching report_request.docx's
    # own style), never "£1,234", to keep the rendered image legible and portable across OSes.
    font = ImageFont.load_default(size=16)
    draw.text((10, 10), title, fill="black", font=font)

    def draw_row(y: int, cells: tuple[str, ...], *, header_row: bool) -> None:
        x = 10
        if header_row:
            draw.rectangle([x, y, x + sum(col_widths), y + row_height], fill=(230, 230, 230))
        for cell, w in zip(cells, col_widths, strict=True):
            draw.rectangle([x, y, x + w, y + row_height], outline="black")
            draw.text((x + 5, y + row_height // 2 - 5), cell, fill="black", font=font)
            x += w

    draw_row(top, header, header_row=True)
    for i, row in enumerate(rows):
        draw_row(top + row_height * (i + 1), row, header_row=False)
    # compress_level=0: the pixels are identical on every platform, but zlib's compressed output is
    # not (macOS and the Linux CI runner gave different bytes for the same image), and the rebuild
    # test compares bytes. Stored (uncompressed) blocks are byte-identical everywhere.
    img.save(path, format="PNG", compress_level=0, optimize=False)


def build_case(src_path: Path, out_root: Path = OUT_DIR) -> Path:
    case_id = src_path.stem  # "case_01"
    markdown = src_path.read_text(encoding="utf-8")
    sections = _split_sections(markdown)
    out_dir = out_root / case_id
    out_dir.mkdir(parents=True, exist_ok=True)

    meeting_n = 0
    instruction_n = 0
    unknown_n = 0

    for section in sections:
        heading = section.heading
        if heading == "Account data":
            m = JSON_BLOCK_RE.search(section.body)
            if m is None:
                raise ValueError(f"{src_path}: 'Account data' has no ```json block")
            data = json.loads(m.group(1))
            (out_dir / "client_data_db.json").write_text(
                json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
            )

        elif heading == "Meeting notes":
            meeting_n += 1
            default_name = (
                "meeting_notes.docx" if meeting_n == 1 else f"meeting_notes_{meeting_n}.docx"
            )
            filename = section.directives.get("Filename", default_name)
            paragraphs = _paragraphs(section.body)
            _write_paragraphs_docx(out_dir / filename, paragraphs, _docx_date(section))

        elif heading == "Report instruction":
            instruction_n += 1
            default_name = (
                "report_request.docx"
                if instruction_n == 1
                else f"report_request_{instruction_n}.docx"
            )
            filename = section.directives.get("Filename", default_name)
            fields = [(r[0], r[1]) for r in _table_rows(section.body) if len(r) == 2]
            _write_instruction_docx(out_dir / filename, fields, _docx_date(section))

        elif heading == "Statement image":
            filename = section.directives.get("Filename", "statement_summary.png")
            title = section.directives.get("Title", "Holloway Account Statement")
            rows_text = section.body
            m = re.search(r"^.*\|.*\|.*\|.*$\n(.*)", rows_text, re.MULTILINE | re.DOTALL)
            data_rows = _table_rows(m.group(1)) if m else []
            _render_statement_image(out_dir / filename, title, data_rows)  # type: ignore[arg-type]

        elif heading == "Internal guidance":
            filename = section.directives.get("Filename", "fde_notes.md")
            text = (
                "# Internal notes: data sources\n\n"
                "For whoever configures the report: what each source is, and how they relate.\n\n"
                "## This client\n"
                f"{section.body.strip()}\n"
            )
            (out_dir / filename).write_text(text, encoding="utf-8")

        elif heading == "Unknown document":
            unknown_n += 1
            default_name = "welcome_pack.docx" if unknown_n == 1 else f"unknown_{unknown_n}.docx"
            filename = section.directives.get("Filename", default_name)
            paragraphs = _paragraphs(section.body)
            _write_paragraphs_docx(out_dir / filename, paragraphs, _docx_date(section))

        # "Rule" / "Expected release" are metadata lines, not "## " sections, so nothing to build.

    return out_dir


def main(argv: list[str]) -> int:
    out_root = OUT_DIR
    if "--out-dir" in argv:
        i = argv.index("--out-dir")
        out_root = Path(argv[i + 1])
        argv = argv[:i] + argv[i + 2 :]
    names = argv or sorted(p.stem for p in SRC_DIR.glob("case_*.md"))
    for name in names:
        src_path = SRC_DIR / f"{name}.md"
        if not src_path.exists():
            print(f"no such case source: {src_path}", file=sys.stderr)
            return 2
        out_dir = build_case(src_path, out_root)
        print(f"built {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
